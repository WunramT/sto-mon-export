"""Phase 7: Export (D19, Golden-File byteweise), Review-Blatt, Regression (D23), Views und Checks."""

import os
from pathlib import Path

import pytest
import sqlalchemy as sa
from openpyxl import load_workbook

from basis_bom import checks, export, lauf, loader, regress, review, rules
from basis_bom.explode import Aufloeser
from basis_bom.source import SapSource

GOLDEN = Path(__file__).parent / "golden"


@pytest.fixture(scope="module")
def fixture_src():
    return SapSource.from_ladeergebnis(loader.lade_fixtures())


def test_export_golden(fixture_src, tmp_path):
    erg = Aufloeser(fixture_src, rules.seed_regelstand()).loese_alle(["90000001"])
    pfade = export.exportiere(erg.df(), fixture_src, tmp_path)
    neu = pfade[0].read_bytes()
    golden = GOLDEN / "grundversion_sap_format.csv"
    if os.environ.get("BASIS_BOM_GOLDEN_UPDATE") == "1":
        golden.write_bytes(neu)
    assert neu == golden.read_bytes()
    assert (tmp_path / "90000001" / "90000001_sap_format.csv").read_bytes() == neu
    kopf, root, *pos = neu.decode().splitlines()
    assert kopf == "Werk;Material;ObjektId;Materialkurztext DE;Menge;ME;PTp;Disp.;Warengrp;MArt;SoB"
    assert root.startswith("4000;90000001;;Sessel STO Basis;;;;")
    assert len(pos) == 10  # nur basis/unbedingt


def _ausfuellen(pfad: Path, urteil: str = "richtig") -> None:
    wb = load_workbook(pfad)
    ws = wb[review.BLATT]
    spalte = review.SPALTEN.index("Urteil") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(r, spalte).value = urteil
    wb.save(pfad)


def test_review_regression_wird_rot(fixture_db, tmp_path):
    eng = fixture_db
    lid, erg = lauf.fuehre_aufloesung_aus(eng)
    df = lauf.lade_aufloesung(eng, lid)
    assert len(df) == len(erg.zeilen) and isinstance(df.iloc[0]["spur"], dict)
    blatt = review.review_blatt(df, SapSource.from_db(eng), lid, "90000001", lauf.lade_statistik(eng, lid),
                                tmp_path / "review.xlsx")  # fmt: skip
    ws = load_workbook(blatt)[review.BLATT]
    assert ws.max_row - 1 == len(df)
    assert ws.row_dimensions[ws.max_row].hidden  # ausgeschlossene am Ende eingeklappt

    _ausfuellen(blatt)
    imp = review.importiere(eng, blatt, "test")
    assert imp["bestaetigt"] == ["90000001"] and imp["ohne_urteil"] == 0

    assert regress.regress(eng, lid).empty
    assert lauf.lauf_status(eng, lid) == "regression_ok"

    # absichtlich geänderte Regel: FUNKTION X nicht mehr Basis → MANUEL gewinnt auf Ebene 1
    rules.setze_regel(eng, "FUNKTION", "X", "NICHT_BASIS", geaendert_von="test")
    try:
        lid2, _ = lauf.fuehre_aufloesung_aus(eng)
        abw = regress.regress(eng, lid2)
        assert not abw.empty and set(abw["art"]) >= {"fehlt_im_lauf", "zusaetzlich_im_lauf"}
        assert lauf.lauf_status(eng, lid2) == "regression_fehlgeschlagen"
    finally:
        rules.setze_regel(eng, "FUNKTION", "X", "BASIS", rang=1, geaendert_von="test")


def test_review_import_ungueltiges_urteil(fixture_db, tmp_path):
    lid, _ = lauf.fuehre_aufloesung_aus(fixture_db)
    df = lauf.lade_aufloesung(fixture_db, lid)
    blatt = review.review_blatt(df, SapSource.from_db(fixture_db), lid, "90000001", {}, tmp_path / "r.xlsx")
    _ausfuellen(blatt, "vielleicht")
    with pytest.raises(ValueError, match="unbekannt"):
        review.importiere(fixture_db, blatt, "test")


def test_views(fixture_db):
    lauf.fuehre_aufloesung_aus(fixture_db)
    with fixture_db.connect() as con:

        def q(sql):
            return con.execute(sa.text(sql)).mappings().all()

        ursachen = {r["ursache"] for r in q("SELECT * FROM basis_bom.offene_faelle_ursache")}
        assert {"kein_rang", "nicht_parsbar", "klassenposition", "unbekannter_teilwert",
                "ausgeschlossen_neben_offen"} <= ursachen  # fmt: skip
        assert {r["merkmal"] for r in q("SELECT * FROM basis_bom.kein_rang")} >= {
            "SITZQUALI",
            "PP4000_KS_VERERBEN",
        }
        assert q("SELECT * FROM basis_bom.prozeduren")[0]["prozedur"] == "PP_PREISFINDUNG"
        stat = q("SELECT * FROM basis_bom.statistik_lauf LIMIT 1")[0]
        assert stat["basis"] == 5 and stat["roots_uebersprungen"] == 3
        abdeckung = {
            (r["merkmal"], r["wert"]): r["positionen"] for r in q("SELECT * FROM basis_bom.regel_abdeckung")
        }
        assert abdeckung[("SITZQUALI", "FK")] >= 2 and abdeckung[("FUNKTION", "WA1")] == 0


def test_checks(fixture_db):
    df = checks.pruefe(fixture_db)
    assert len(df) > 10
    assert df.loc[df["pruefung"] == "export: STLNR in STKO", "ok"].item() is False  # 1004 ohne STKO
    assert df.loc[df["datei"] == "regel_002_cabn.sql", "ok"].isna().all()  # CABN fehlt → übersprungen
