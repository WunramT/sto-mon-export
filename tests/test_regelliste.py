"""Regel-Arbeitsliste: Export mit Vorkommen, Import mit Prüfung und Rangtausch."""

from openpyxl import load_workbook
from typer.testing import CliRunner

from basis_bom import lauf, regelliste, rules
from basis_bom.cli import app


def _setze(pfad, blatt, spalten, suche: dict, werte: dict):
    wb = load_workbook(pfad)
    ws = wb[blatt]
    kopf = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        if all(ws.cell(r, kopf.index(k) + 1).value == v for k, v in suche.items()):
            for k, v in werte.items():
                ws.cell(r, kopf.index(k) + 1).value = v
    wb.save(pfad)


def test_export_import(fixture_db, tmp_path, monkeypatch):
    eng = fixture_db
    lauf.fuehre_aufloesung_aus(eng)
    ziel = regelliste.exportiere(eng, tmp_path / "liste.xlsx")
    regeln, kuerzel = regelliste.arbeitsliste(eng)
    oben = regeln.iloc[0]
    assert oben["Status"] == "OFFEN" and oben["Positionen"] > 0  # vorkommende OFFEN-Werte zuerst
    assert regeln.set_index(["Merkmal", "Wert"]).loc[("SITZQUALI", "HR"), "Positionen"] >= 3
    assert "ZZ" in set(kuerzel["Kürzel"])

    # Rangtausch X ↔ MANUEL, FK wird BASIS 2, Kürzel ZZ → SITZTIEFE
    _setze(ziel, "Regeln", regelliste.R_SPALTEN, {"Merkmal": "FUNKTION", "Wert": "X"}, {"Neuer Rang": 2})
    _setze(ziel, "Regeln", regelliste.R_SPALTEN, {"Merkmal": "FUNKTION", "Wert": "MANUEL"}, {"Neuer Rang": 1})
    _setze(ziel, "Regeln", regelliste.R_SPALTEN, {"Merkmal": "SITZQUALI", "Wert": "FK"},
           {"Neuer Status": "BASIS", "Neuer Rang": 2, "Begründung (neu)": "Test"})  # fmt: skip
    _setze(ziel, "Kürzel", regelliste.K_SPALTEN, {"Kürzel": "ZZ"}, {"Neues Merkmal": "SITZTIEFE"})
    monkeypatch.setenv("DATABASE_URL", eng.url.render_as_string(hide_password=False))
    res = CliRunner().invoke(app, ["regel", "import", str(ziel), "--von", "test"])
    assert res.exit_code == 0, res.output
    rs = rules.lade_regelstand(eng)
    assert rs.regel("FUNKTION", "MANUEL").rang == 1 and rs.regel("FUNKTION", "X").rang == 2
    assert rs.regel("SITZQUALI", "FK") == rules.Regel("SITZQUALI", "FK", "BASIS", 2)
    assert rs.aliasse["ZZ"] == rules.Alias("ZZ", "SITZTIEFE", "BASIS")
    # zurücksetzen für andere Tests
    rules.setze_regel(eng, "SITZQUALI", "FK", "OFFEN", geaendert_von="test")
    regelliste.uebernehme(eng, [regelliste.Aenderung("FUNKTION", "X", "BASIS", 1, None),
                                regelliste.Aenderung("FUNKTION", "MANUEL", "BASIS", 2, None)], [], "test")  # fmt: skip
    rules.setze_alias(eng, "ZZ", None, "OFFEN", "test")


def test_import_prueft_doppelten_rang(fixture_db, tmp_path, monkeypatch):
    eng = fixture_db
    lauf.fuehre_aufloesung_aus(eng)
    ziel = regelliste.exportiere(eng, tmp_path / "liste.xlsx")
    _setze(ziel, "Regeln", regelliste.R_SPALTEN, {"Merkmal": "SITZQUALI", "Wert": "FK"}, {"Neuer Rang": 1})
    _setze(
        ziel,
        "Regeln",
        regelliste.R_SPALTEN,
        {"Merkmal": "SITZQUALI", "Wert": "BS"},
        {"Neuer Status": "BASIS"},
    )
    monkeypatch.setenv("DATABASE_URL", eng.url.render_as_string(hide_password=False))
    res = CliRunner().invoke(app, ["regel", "import", str(ziel), "--von", "test"])
    assert res.exit_code == 1
    assert "Rang 1 doppelt" in res.output and "BASIS braucht einen Rang" in res.output
    assert rules.lade_regelstand(eng).status("SITZQUALI", "FK") == "OFFEN"  # nichts geschrieben
