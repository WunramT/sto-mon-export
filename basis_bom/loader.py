"""Laden der SAP-Exporte nach `sap_raw` (AGENT-PLAN Phase 2, EXPORT-PLAN Phase 1 in Python).

Reihenfolge ist zwingend, weil die Schlüsselmengen aufeinander aufbauen (EXPORT-PLAN Phase 1):
MAST → S → STKO/STAS/STPO (chunkweise) → P, K → CUOB → W → CUKB/CUKBT → MARC → M → MARA/MAKT → CABN.
Gefiltert wird beim Laden, nicht danach. D14 (Gültigkeit) wendet erst `source.py` an; hier bleiben
gelöschte und zukünftige Zeilen erhalten.
"""

from __future__ import annotations

import ast
import csv
import datetime as dt
import io
import logging
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
import sqlalchemy as sa
from sqlalchemy.engine import Engine

from . import config, headers

log = logging.getLogger(__name__)

# docs/EXPORTE.md „Maßgebliche Dateien“. Alles andere im Export-Ordner wird ignoriert.
MASSGEBLICH: dict[str, str] = {
    "MAST": "EXPORT_mast_20260923_150612.XLSX",
    "STPO": "fa_stpo_202606291123.csv",
    "STAS": "EXPORT_stas_20260923_142938.XLSX",
    "STKO": "EXPORT_stko_20260923_144510.XLSX",
    "CUOB": "EXPORT_20260512_153230_cuob.XLSX",
    "CUKB": "EXPORT_20260512_153418_CUKB.XLSX",
    "CUKBT": "EXPORT_CUKBT_20260923_150301.XLSX",
    "MARA": "EXPORT_mara_20260923_145220.XLSX",
    "MARC": "EXPORT_marc_20260923_141650.XLSX",
    "MAKT": "EXPORT_makt_20260923_144656.XLSX",
    "CABN": "EXPORT_cabn_20260923_142303.XLSX",
}
ROOT_XLSX = "Kopie von 4000_STO_OKU_Alle Modelle_Planzeiten_Merkmal.xlsx"
SCHLUESSEL_DATEIEN = {"S": "S.csv", "P_idnrk": "P_idnrk.csv", "M": "M.csv", "W": "W.csv"}

LADE_REIHENFOLGE = ["MAST", "STKO", "STAS", "STPO", "CUOB", "CUKB", "CUKBT", "MARC", "MARA", "MAKT", "CABN",
                    "CAWN", "CAWNT"]  # fmt: skip

# Führende Nullen einheitlich entfernen (Schlüssel und Zähler).
NULLEN_WEG = {"MATNR", "IDNRK", "STLNR", "KNOBJ", "KNNUM", "STLAL", "STLKN", "STPOZ", "STASZ", "STKOZ",
              "ADZHL", "KNSRT", "STLAN", "ATINN", "ATZHL"}  # fmt: skip
DATUMSSPALTEN = {"DATUV"}
ZAHLSPALTEN = {"MENGE", "BMENG"}

# Kanonische Merkmale (EXPORT-PLAN Phase 3), falls die Alias-Tabelle noch nicht geladen ist.
MERKMALLISTE_DEFAULT = ["RUECKEN_OPTIK", "RUECKEN_FUNK", "ARM_OPTIK", "ARM_L", "ARM_R", "SITZQUALI", "SITZHOEHE",
                        "SITZTIEFE", "FUNKTION", "ELEKTRO", "GASDRUCK", "FUSS", "3_FUSS", "AKKU", "MOTOR"]  # fmt: skip

EXCEL_MAX_ZEILEN = 1_048_576
STPO_CHUNK = 200_000


class HeaderFehler(ValueError):
    """Pflichtspalte nicht zugeordnet; trägt das bisherige Ladeergebnis für den Header-Bericht."""

    def __init__(self, text: str, erg: Ladeergebnis | None = None) -> None:
        super().__init__(text)
        self.erg = erg


@dataclass
class Quelle:
    tabelle: str
    pfad: Path
    export_datum: dt.date | None


@dataclass
class Protokoll:
    tabelle: str
    datei: str
    export_datum: dt.date | None
    zeilen_roh: int = 0
    zeilen_geladen: int = 0
    spalten: list[str] = field(default_factory=list)
    unbekannte_spalten: list[str] = field(default_factory=list)
    fehlende_spalten: list[str] = field(default_factory=list)
    trenner: str | None = None
    blaetter: dict[str, int] = field(default_factory=dict)
    fuehrende_nullen: bool | None = None
    hinweise: list[str] = field(default_factory=list)
    header: dict[str, str | None] = field(
        default_factory=dict
    )  # Original-Header → technisch (None = unbekannt)


@dataclass
class Ladeergebnis:
    tabellen: dict[str, pd.DataFrame] = field(default_factory=dict)
    protokoll: dict[str, Protokoll] = field(default_factory=dict)
    schluessel: dict[str, set[str]] = field(default_factory=dict)
    roots: pd.DataFrame | None = None
    root_info: dict = field(default_factory=dict)
    root_ausschluss: pd.DataFrame | None = None
    abgleich: dict[str, dict] = field(default_factory=dict)

    @property
    def export_daten(self) -> dict[str, dt.date]:
        return {t: p.export_datum for t, p in self.protokoll.items() if p.export_datum}


# ---------------------------------------------------------------------------------------------------------
# Werte normalisieren


def export_datum_aus_name(name: str) -> dt.date | None:
    m = re.search(r"(20\d{2})(\d{2})(\d{2})", name)
    if not m:
        return None
    try:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def ohne_nullen(v: str) -> str:
    s = str(v).strip()
    if re.fullmatch(r"\d+\.0+", s):  # Excel-Zahl „11071032.0“
        s = s.split(".")[0]
    if s.isdigit():
        return s.lstrip("0") or "0"
    return s


def parse_datum(v: object) -> dt.date | None:
    if v is None:
        return None
    s = str(v).strip()
    if not s or s in {"00000000", "0", "nan", "NaT", "None"}:
        return None
    for pat, fmt in (
        (r"\d{8}", "%Y%m%d"),
        (r"\d{4}-\d{2}-\d{2}", "%Y-%m-%d"),
        (r"\d{2}\.\d{2}\.\d{4}", "%d.%m.%Y"),
    ):
        m = re.match(pat, s)
        if m:
            try:
                return dt.datetime.strptime(m.group(0), fmt).date()
            except ValueError:
                return None
    return None


def parse_zahl(v: object) -> float | None:
    s = str(v).strip() if v is not None else ""
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


# ---------------------------------------------------------------------------------------------------------
# Dateien lesen


def _trenner(pfad: Path) -> tuple[str, str]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            with pfad.open(encoding=enc) as f:
                kopf = f.readline()
            break
        except UnicodeDecodeError:
            continue
    sep = ";" if kopf.count(";") >= kopf.count(",") else ","
    return sep, enc


def lese_roh(pfad: Path, prot: Protokoll, chunksize: int | None = None) -> Iterator[pd.DataFrame]:
    """Rohdaten als str-DataFrames; CSV optional chunkweise."""
    if pfad.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        blaetter = pd.read_excel(pfad, sheet_name=None, dtype=str, keep_default_na=False)
        for name, df in blaetter.items():
            prot.blaetter[name] = len(df) + 1  # + Kopfzeile
            if len(df) + 1 >= EXCEL_MAX_ZEILEN:
                prot.hinweise.append(
                    f"Blatt {name!r} hat {len(df) + 1} Zeilen – Excel-Grenze, evtl. abgeschnitten"
                )
        dfs = [df for df in blaetter.values() if len(df.columns)]
        if dfs:
            yield pd.concat(dfs, ignore_index=True) if len(dfs) > 1 else dfs[0]
        return
    sep, enc = _trenner(pfad)
    prot.trenner = sep
    reader = pd.read_csv(pfad, sep=sep, dtype=str, keep_default_na=False, encoding=enc, chunksize=chunksize)
    if chunksize is None:
        yield reader
    else:
        yield from reader


def normalisiere(df: pd.DataFrame, tabelle: str, prot: Protokoll) -> pd.DataFrame:
    """Header technisch benennen, Werte trimmen, Schlüssel ohne führende Nullen, Datum/Zahl parsen."""
    namen: dict[str, str] = {}
    for col in df.columns:
        tech = headers.technischer_name(tabelle, col)
        if tech is None:
            tech = headers.ersatzname(col)
            if tech not in prot.unbekannte_spalten:
                prot.unbekannte_spalten.append(tech)
                log.debug("%s: unbekannte Spalte %r → %s (wird mitgeladen)", tabelle, col, tech)
        prot.header.setdefault(str(col), headers.technischer_name(tabelle, col))
        while tech in namen.values():
            tech += "_2"
        namen[col] = tech
    df = df.rename(columns=namen)
    if not prot.spalten:
        prot.spalten = list(df.columns)
        prot.fehlende_spalten = [c for c in headers.SPALTEN.get(tabelle, []) if c not in df.columns]
        zugeordnet = {
            str(o): t for o, t in namen.items() if t in headers.SPALTEN.get(tabelle, []) and str(o) != t
        }
        if zugeordnet:
            log.info("%s: Header zugeordnet %s", tabelle, zugeordnet)
        if prot.unbekannte_spalten:
            log.info("%s: %s unbekannte Spalten werden mitgeladen (Details mit -v)", tabelle,
                     len(prot.unbekannte_spalten))  # fmt: skip
        if prot.fehlende_spalten:
            log.warning("%s: Spalten fehlen: %s", tabelle, prot.fehlende_spalten)
    for col in df.columns:
        df[col] = df[col].astype(str).str.strip()
    for col in NULLEN_WEG & set(df.columns):
        if prot.fuehrende_nullen is None and col in {"MATNR", "IDNRK", "STLNR"}:
            prot.fuehrende_nullen = bool(df[col].str.match(r"^0\d").any())
        df[col] = ohne_nullen_serie(df[col])
    for col in DATUMSSPALTEN & set(df.columns):
        df[col] = datum_serie(df[col])
    for col in ZAHLSPALTEN & set(df.columns):
        df[col] = zahl_serie(df[col])
    return df


# Vektorisierte Fassungen von ohne_nullen/parse_datum/parse_zahl (STPO-CSV hat ~1 GB).


def ohne_nullen_serie(s: pd.Series) -> pd.Series:
    s = s.str.replace(r"^(\d+)\.0+$", r"\1", regex=True)
    ziffern = s.str.fullmatch(r"\d+").fillna(False).astype(bool)
    return s.where(~ziffern, s.str.lstrip("0").replace("", "0"))


def datum_serie(s: pd.Series) -> pd.Series:
    s = s.astype(str).str.strip()
    ergebnis = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    for muster, laenge, fmt in ((r"\d{8}", 8, "%Y%m%d"), (r"\d{4}-\d{2}-\d{2}", 10, "%Y-%m-%d"),
                                (r"\d{2}\.\d{2}\.\d{4}", 10, "%d.%m.%Y")):  # fmt: skip
        maske = ergebnis.isna() & s.str.match(muster).fillna(False).astype(bool)
        if maske.any():
            ergebnis[maske] = pd.to_datetime(s[maske].str[:laenge], format=fmt, errors="coerce")
    werte = ergebnis.dt.date.astype(object)
    werte[ergebnis.isna()] = None
    return werte


def zahl_serie(s: pd.Series) -> pd.Series:
    s = s.astype(str).str.strip()
    komma = s.str.contains(",", regex=False)
    s = s.where(~komma, s.str.replace(".", "", regex=False).str.replace(",", ".", regex=False))
    return pd.to_numeric(s.replace("", None), errors="coerce")


# ---------------------------------------------------------------------------------------------------------
# Pipeline


def _filter(df: pd.DataFrame, **bedingungen: Iterable[str] | str) -> pd.DataFrame:
    maske = pd.Series(True, index=df.index)
    for col, werte in bedingungen.items():
        if col not in df.columns:
            log.warning("Filterspalte %s fehlt – Filter entfällt (Header-Mapping prüfen)", col)
            continue
        if isinstance(werte, str):
            maske &= df[col] == werte
        else:
            maske &= df[col].isin(set(werte))
    return df[maske]


def _leer(v: pd.Series) -> pd.Series:
    return v.isna() | (v.astype(str).str.strip() == "")


def lade_quellen(
    quellen: dict[str, Quelle],
    merkmalliste: Iterable[str] | None = None,
    root_materialien: Iterable[str] = (),
) -> Ladeergebnis:
    """Liest alle Quellen, filtert über die Schlüsselmengen und liefert DataFrames (ohne DB)."""
    erg = Ladeergebnis()
    werks, stlan = config.WERKS, config.STLAN

    def lies(tab: str, chunksize: int | None = None, filt=None, vorfilter=None) -> pd.DataFrame | None:
        q = quellen.get(tab)
        if q is None:
            log.warning("%s: keine Quelle – Ersatzregel laut EXPORT-PLAN „Übergang“", tab)
            return None
        prot = Protokoll(tab, q.pfad.name, q.export_datum)
        erg.protokoll[tab] = prot
        teile = []
        for roh in lese_roh(q.pfad, prot, chunksize):
            prot.zeilen_roh += len(roh)
            if vorfilter:  # Schlüsselspalte vor der vollen Normalisierung filtern (Laufzeit bei STPO)
                spalte, werte = vorfilter
                roh_col = next((c for c in roh.columns if headers.technischer_name(tab, c) == spalte), None)
                if roh_col is not None:
                    roh = roh[ohne_nullen_serie(roh[roh_col].astype(str).str.strip()).isin(werte)]
            df = normalisiere(roh, tab, prot)
            teile.append(filt(df) if filt else df)
        df = pd.concat(teile, ignore_index=True) if teile else pd.DataFrame(columns=prot.spalten)
        prot.zeilen_geladen = len(df)
        log.info("%s: %s → %s Zeilen (%s)", tab, prot.zeilen_roh, prot.zeilen_geladen, q.pfad.name)
        erg.tabellen[tab] = df
        return df

    # MAST → S
    mast = lies("MAST", filt=lambda d: _filter(d, WERKS=werks, STLAN=stlan))
    if mast is None:
        raise ValueError("MAST fehlt – ohne MAST keine Schlüsselmenge S")
    fehlend = [c for c in ("MATNR", "WERKS", "STLAN", "STLNR") if c not in mast.columns]
    if fehlend:
        unbekannt = [h for h, t in erg.protokoll["MAST"].header.items() if t is None]
        raise HeaderFehler(
            f"MAST: Pflichtspalten {fehlend} nicht zugeordnet – ohne sie ist S falsch (Verwendung 1). "
            f"Unbekannte MAST-Header: {unbekannt}. Zuordnung in docs/header_mapping.csv eintragen "
            f"(z. B. 'MAST;<Header>;STLAN') oder `basis-bom db headers --from-dir <pfad>` ansehen.",
            erg,
        )
    s_mast = mast[_leer(mast["LKENZ"])] if "LKENZ" in mast.columns else mast
    S = set(s_mast["STLNR"])
    erg.schluessel["S"] = S

    lies("STKO", filt=lambda d: _filter(d, STLNR=S))
    lies("STAS", filt=lambda d: _filter(d, STLNR=S))
    stpo = lies("STPO", chunksize=STPO_CHUNK, filt=lambda d: _filter(d, STLNR=S), vorfilter=("STLNR", S))
    if stpo is None:
        raise ValueError("STPO fehlt")
    erg.schluessel["P_idnrk"] = set(stpo["IDNRK"]) - {""}
    K = set(stpo["KNOBJ"]) - {"", "0"}
    erg.schluessel["K"] = K

    cuob = lies("CUOB", filt=lambda d: _filter(d, KNTAB="STPO", KNOBJ=K))
    if cuob is None:
        raise ValueError("CUOB fehlt")
    W = set(cuob["KNNUM"]) - {""}
    erg.schluessel["W"] = W
    lies("CUKB", filt=lambda d: _filter(d, KNNUM=W))
    lies("CUKBT", filt=lambda d: _filter(_filter(d, KNNUM=W), SPRAS=["D", "DE"]))

    marc = lies("MARC", filt=lambda d: _filter(d, WERKS=werks))
    M = set(erg.schluessel["P_idnrk"]) | set(mast["MATNR"]) | {ohne_nullen(m) for m in root_materialien}
    if marc is not None:
        M |= set(marc["MATNR"])
    erg.schluessel["M"] = M
    lies("MARA", filt=lambda d: _filter(d, MATNR=M))
    lies("MAKT", filt=lambda d: _filter(_filter(d, MATNR=M), SPRAS=["D", "DE"]))

    merkmale = {m.upper() for m in (merkmalliste or MERKMALLISTE_DEFAULT)}
    cabn = lies("CABN", filt=lambda d: _filter(d, ATNAM=merkmale) if "ATNAM" in d.columns else d.iloc[0:0])
    if cabn is not None and "ATINN" not in cabn.columns:
        log.warning("CABN: Spalte ATINN fehlt – CAWN/CAWNT nicht filterbar, Marker cawn_fehlt bleibt")
    elif cabn is not None:
        A = set(cabn["ATINN"])
        erg.schluessel["A"] = A
        lies("CAWN", filt=lambda d: _filter(d, ATINN=A))
        lies("CAWNT", filt=lambda d: _filter(_filter(d, ATINN=A), SPRAS=["D", "DE"]))
    return erg


def fixture_quellen(verz: Path) -> dict[str, Quelle]:
    """`tests/fixtures/<tabelle>_<JJJJMMTT>.csv`."""
    quellen = {}
    for p in sorted(verz.glob("*.csv")):
        m = re.fullmatch(r"([a-z]+)_(\d{8})", p.stem)
        if m and m.group(1).upper() in headers.SPALTEN:
            quellen[m.group(1).upper()] = Quelle(m.group(1).upper(), p, export_datum_aus_name(p.stem))
    return quellen


def verzeichnis_quellen(verz: Path) -> dict[str, Quelle]:
    """Nur die in docs/EXPORTE.md als maßgeblich genannten Dateien (Groß-/Kleinschreibung egal)."""
    vorhanden = {p.name.lower(): p for p in verz.iterdir() if p.is_file()} if verz.is_dir() else {}
    quellen = {}
    for tab, name in MASSGEBLICH.items():
        p = vorhanden.get(name.lower())
        if p is None:
            log.warning("%s: maßgebliche Datei %s fehlt in %s", tab, name, verz)
            continue
        quellen[tab] = Quelle(tab, p, export_datum_aus_name(name))
    return quellen


# ---------------------------------------------------------------------------------------------------------
# Root-Materialien (D22)


def lese_root_csv(pfad: Path) -> pd.DataFrame:
    df = pd.read_csv(pfad, dtype=str, keep_default_na=False)
    return pd.DataFrame({"matnr": sorted({ohne_nullen(m) for m in df["matnr"] if m.strip()})})


def lese_root_xlsx(pfad: Path) -> tuple[pd.DataFrame, dict]:
    """Planzeiten-XLSX: Spalte mit den meisten Materialnummern über alle Blätter; Rest bleibt ungeladen."""
    blaetter = pd.read_excel(pfad, sheet_name=None, dtype=str, header=None, keep_default_na=False)
    info: dict = {"blaetter": {n: len(df) for n, df in blaetter.items()}}
    bester = (0, None, None)
    for name, df in blaetter.items():
        for col in df.columns:
            werte = df[col].astype(str).str.strip().map(ohne_nullen)
            treffer = int(werte.str.fullmatch(r"\d{7,10}").sum())
            if treffer > bester[0]:
                bester = (treffer, name, col)
    if bester[1] is None:
        raise ValueError(f"{pfad.name}: keine Spalte mit Materialnummern gefunden")
    df = blaetter[bester[1]]
    werte = df[bester[2]].astype(str).str.strip().map(ohne_nullen)
    werte = werte[werte.str.fullmatch(r"\d{7,10}")]
    kopf = [str(v) for v in df[bester[2]].head(5) if not re.fullmatch(r"\d+(\.0)?", str(v).strip())]
    info.update(
        blatt=bester[1],
        spalte=int(bester[2]) + 1,
        spaltenkopf=kopf[0] if kopf else None,
        werte=len(werte),
        duplikate=int(werte.duplicated().sum()),
        eindeutig=werte.nunique(),
    )
    return pd.DataFrame({"matnr": sorted(set(werte))}), info


def legacy_material_list(pfad: Path = config.LEGACY_SCRIPT) -> list[str]:
    """`MATERIAL_LIST` aus legacy/basis_bom_v0.py, ohne das Skript auszuführen."""
    tree = ast.parse(pfad.read_text(encoding="utf-8"))
    for node in ast.walk(tree):
        if isinstance(node, ast.Assign) and any(
            getattr(t, "id", None) == "MATERIAL_LIST" for t in node.targets
        ):
            return [ohne_nullen(v) for v in ast.literal_eval(node.value)]
    return []


def schluessel_abgleich(erg: Ladeergebnis, verz: Path) -> None:
    """S/P/M/W gegen die manuell erzeugten Dateien (docs/EXPORTE.md) – nur Protokoll."""
    for name, datei in SCHLUESSEL_DATEIEN.items():
        p = verz / datei
        if not p.exists() or name not in erg.schluessel:
            continue
        sep, enc = _trenner(p)
        ref_df = pd.read_csv(p, sep=sep, dtype=str, keep_default_na=False, encoding=enc)
        ref = {ohne_nullen(v) for v in ref_df.iloc[:, 0] if str(v).strip()}
        ist = erg.schluessel[name]
        erg.abgleich[name] = {
            "datei": len(ref), "berechnet": len(ist),
            "nur_datei": sorted(ref - ist)[:20], "nur_berechnet": sorted(ist - ref)[:20],
            "anzahl_nur_datei": len(ref - ist), "anzahl_nur_berechnet": len(ist - ref),
        }  # fmt: skip
        log.info("Abgleich %s: Datei %s, berechnet %s, nur Datei %s, nur berechnet %s", name, len(ref), len(ist),
                 len(ref - ist), len(ist - ref))  # fmt: skip


# ---------------------------------------------------------------------------------------------------------
# Schreiben nach Postgres


def _leer_wert(v) -> bool:
    """None, NaN, NaT, pd.NA → SQL NULL."""
    if v is None or isinstance(v, str):
        return v is None
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def _sql_typ(col: str) -> str:
    if col in DATUMSSPALTEN or col == "EXPORT_DATUM":
        return "date"
    if col in ZAHLSPALTEN:
        return "numeric"
    return "text"


def schreibe_tabelle(eng: Engine, tabelle: str, df: pd.DataFrame, export_datum: dt.date | None) -> None:
    """Ersetzt `sap_raw.<tabelle>` per COPY. Spalten klein geschrieben, plus `export_datum`."""
    df = df.copy()
    df["EXPORT_DATUM"] = export_datum
    cols = list(df.columns)
    name = f"sap_raw.{tabelle.lower()}"
    ddl = ", ".join(f'"{c.lower()}" {_sql_typ(c)}' for c in cols)
    raw = eng.raw_connection()
    try:
        with raw.cursor() as cur:
            cur.execute(f"DROP TABLE IF EXISTS {name} CASCADE")
            cur.execute(f"CREATE TABLE {name} ({ddl})")
            buf = io.StringIO()
            w = csv.writer(buf, lineterminator="\n")
            for row in df.itertuples(index=False, name=None):
                w.writerow(["\\N" if _leer_wert(v) else v for v in row])
            buf.seek(0)
            collist = ", ".join(f'"{c.lower()}"' for c in cols)
            with cur.copy(f"COPY {name} ({collist}) FROM STDIN WITH (FORMAT csv, NULL '\\N')") as cp:
                while chunk := buf.read(1 << 20):
                    cp.write(chunk)
            for key in ("stlnr", "matnr", "knobj", "knnum", "idnrk"):
                if key in (c.lower() for c in cols):
                    cur.execute(f"CREATE INDEX ON {name} ({key})")
        raw.commit()
    finally:
        raw.close()


def schreibe(eng: Engine, erg: Ladeergebnis, quelle_roots: str) -> None:
    for tab, df in erg.tabellen.items():
        schreibe_tabelle(eng, tab, df, erg.protokoll[tab].export_datum)
    # nicht gelieferte optionale Tabellen entfernen, damit keine alten Stände liegen bleiben
    with eng.begin() as con:
        for tab in headers.SPALTEN:
            if tab not in erg.tabellen:
                con.execute(sa.text(f"DROP TABLE IF EXISTS sap_raw.{tab.lower()} CASCADE"))
        for p in erg.protokoll.values():
            con.execute(
                sa.text(
                    "INSERT INTO basis_bom.ladeprotokoll (tabelle, datei, export_datum, zeilen_roh, zeilen_geladen,"
                    " unbekannte_spalten, fehlende_spalten, hinweise) VALUES (:t, :d, :e, :r, :g, :u, :f, :h)"
                ),
                {
                    "t": p.tabelle,
                    "d": p.datei,
                    "e": p.export_datum,
                    "r": p.zeilen_roh,
                    "g": p.zeilen_geladen,
                    "u": p.unbekannte_spalten,
                    "f": p.fehlende_spalten,
                    "h": p.hinweise,
                },  # fmt: skip
            )
        if erg.roots is not None:
            schreibe_roots(con, erg.roots["matnr"].tolist(), quelle_roots)
        if erg.root_ausschluss is not None:
            for r in erg.root_ausschluss.itertuples():
                con.execute(
                    sa.text(
                        "INSERT INTO basis_bom.root_ausschluss (matnr, grund, geaendert_von) "
                        "SELECT :m, :g, 'fixture' WHERE NOT EXISTS "
                        "(SELECT 1 FROM basis_bom.root_ausschluss WHERE matnr = :m)"
                    ),
                    {"m": ohne_nullen(r.matnr), "g": r.grund},
                )


def schreibe_roots(con, matnrs: list[str], quelle: str) -> None:
    """Historisiert (D21): neue Nummern einfügen, nicht mehr gelieferte derselben Quelle schließen."""
    con.execute(
        sa.text(
            "UPDATE basis_bom.root_material SET gueltig_bis = now() "
            "WHERE quelle = :q AND gueltig_bis IS NULL AND NOT (matnr = ANY(:m))"
        ),
        {"q": quelle, "m": matnrs},
    )
    con.execute(
        sa.text(
            "INSERT INTO basis_bom.root_material (matnr, quelle, geaendert_von) "
            "SELECT m, :q, 'loader' FROM unnest(CAST(:m AS text[])) AS m "
            "WHERE NOT EXISTS (SELECT 1 FROM basis_bom.root_material r "
            "                  WHERE r.matnr = m AND r.gueltig_bis IS NULL)"
        ),
        {"q": quelle, "m": matnrs},
    )


# ---------------------------------------------------------------------------------------------------------
# Einstiegspunkte


def lade_fixtures(verz: Path = config.FIXTURES_DIR) -> Ladeergebnis:
    erg = lade_quellen(fixture_quellen(verz))
    if (verz / "root_material.csv").exists():
        erg.roots = lese_root_csv(verz / "root_material.csv")
    if (verz / "root_ausschluss.csv").exists():
        erg.root_ausschluss = pd.read_csv(verz / "root_ausschluss.csv", dtype=str, keep_default_na=False)
    return erg


def lade_verzeichnis(verz: Path, merkmalliste: Iterable[str] | None = None) -> Ladeergebnis:
    roots, info = None, {}
    root_pfad = next((p for p in verz.iterdir() if p.name.lower() == ROOT_XLSX.lower()), None)
    if root_pfad is not None:
        roots, info = lese_root_xlsx(root_pfad)
    else:
        log.warning("Planzeiten-XLSX %s fehlt – root_material bleibt unverändert", ROOT_XLSX)
    erg = lade_quellen(verzeichnis_quellen(verz), merkmalliste, roots["matnr"] if roots is not None else ())
    erg.roots, erg.root_info = roots, info
    if roots is not None:
        legacy = set(legacy_material_list())
        neu = set(roots["matnr"])
        S_mat = set(erg.tabellen["MAST"]["MATNR"])
        info.update(
            legacy=len(legacy), nur_xlsx=sorted(neu - legacy), nur_legacy=sorted(legacy - neu),
            ohne_stueckliste=sorted(neu - S_mat),
        )  # fmt: skip
        log.info("Planzeiten: %s Materialien, nur XLSX %s, nur MATERIAL_LIST %s, ohne Stückliste %s", len(neu),
                 len(neu - legacy), len(legacy - neu), len(neu - S_mat))  # fmt: skip
    schluessel_abgleich(erg, verz)
    return erg


def header_bericht(protokolle: Iterable[Protokoll]) -> str:
    """Markdown: pro Tabelle alle Original-Header mit Zuordnung – Grundlage für docs/header_mapping.csv."""
    z = ["# Header-Bericht", "", "Unbekannte Header lassen sich in `docs/header_mapping.csv` zuordnen "
         "(`tabelle;header;technisch`).", ""]  # fmt: skip
    for p in protokolle:
        z.append(f"## {p.tabelle} – {p.datei}")
        z.append("")
        fehlend = [c for c in headers.SPALTEN.get(p.tabelle, []) if c not in p.header.values()]
        z.append(f"Nicht zugeordnet (erwartet laut EXPORT-PLAN): {', '.join(fehlend) or '–'}")
        z.append("")
        z.append("| Header | technisch |")
        z.append("|---|---|")
        for h, t in p.header.items():
            z.append(f"| {h} | {t or '**unbekannt**'} |")
        z.append("")
    return "\n".join(z)


def lese_header(verz: Path) -> list[Protokoll]:
    """Nur die Kopfzeilen der maßgeblichen Dateien (schnell, für `basis-bom db headers`)."""
    from openpyxl import load_workbook

    out = []
    for tab, q in verzeichnis_quellen(verz).items():
        prot = Protokoll(tab, q.pfad.name, q.export_datum)
        if q.pfad.suffix.lower() in {".xlsx", ".xlsm"}:
            wb = load_workbook(q.pfad, read_only=True)
            kopf = next(wb.worksheets[0].iter_rows(max_row=1, values_only=True), ())
            wb.close()
        else:
            sep, enc = _trenner(q.pfad)
            with q.pfad.open(encoding=enc) as f:
                kopf = next(csv.reader(f, delimiter=sep), [])
        for h in kopf:
            if h is not None:
                prot.header[str(h)] = headers.technischer_name(tab, h)
        out.append(prot)
    return out
