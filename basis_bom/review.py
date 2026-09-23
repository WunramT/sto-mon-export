"""Review-Blatt pro Root-Material (D24): Export als XLSX für den Fachbereich und Import der Urteile.

Urteile beziehen sich auf den Status der Zeile:
- `richtig`: Status stimmt (drin bei basis/unbedingt, draußen sonst)
- `fehlt`: Material gehört in die Basis-Version, ist aber nicht drin (auch als neue Zeile ergänzbar)
- `gehoert_nicht_rein`: Zeile ist drin, gehört aber nicht in die Basis-Version
Ist jede Zeile eines Root-Materials `richtig`, gilt es als bestätigt (Tabelle `bestaetigt`, D23).
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.engine import Engine

from . import anzeige
from .explode import EXPORT_STATUS
from .source import SapSource

URTEILE = ("richtig", "fehlt", "gehoert_nicht_rein")
BLATT = "Review"
SPALTEN = ["Root", "Ebene", "Parent", "Material", "Kurztext", "Menge kumuliert", "ME", "Status", "Grund", "Urteil",
           "Kommentar", "Pfad", "Lauf"]  # fmt: skip
BREITEN = {"Root": 11, "Ebene": 6, "Parent": 11, "Material": 11, "Kurztext": 30, "Menge kumuliert": 10, "ME": 5,
           "Status": 20, "Grund": 70, "Urteil": 18, "Kommentar": 40}  # fmt: skip
GRAU = PatternFill("solid", fgColor="EEEEEE")
GELB = PatternFill("solid", fgColor="FFF2CC")


def grund_lesbar(z) -> str:
    spur = anzeige.spur_lesbar(z["spur"]) if isinstance(z["spur"], dict) else ""
    teile = [t for t in (z["grund"], spur) if t]
    return " — ".join(dict.fromkeys(teile))


def _reihenfolge(df: pd.DataFrame) -> pd.DataFrame:
    """Enthaltene und zu prüfende Positionen zuerst (Baumreihenfolge), ausgeschlossene/ignorierte am Ende."""
    d = df.copy()
    d["_raus"] = d["status"].isin(anzeige.EINGEKLAPPT)
    d["_pfad"] = d["pfad"].map(anzeige.pfad_sortierung)
    return d.sort_values(["_raus", "_pfad"], kind="stable").drop(columns=["_raus", "_pfad"])


def review_blatt(
    df: pd.DataFrame, src: SapSource, lauf_id: int, root: str, statistik: dict, ziel: Path
) -> Path:
    maktx = src.lookup("MAKT", "MAKTX")
    wb = Workbook()
    ws = wb.active
    ws.title = BLATT
    ws.append(SPALTEN)
    for c in ws[1]:
        c.font = Font(bold=True)
    teil = _reihenfolge(df[df["root_matnr"] == root])
    for _, z in teil.iterrows():
        ws.append([root, int(z["ebene"]), z["parent_matnr"], z["matnr"], maktx.get(z["matnr"], ""),
                   None if pd.isna(z["menge_kum"]) else float(z["menge_kum"]), z["meins"], z["status"],
                   grund_lesbar(z), None, None, z["pfad"], lauf_id])  # fmt: skip
        r = ws.max_row
        if z["status"] in anzeige.EINGEKLAPPT:
            ws.row_dimensions[r].outlineLevel = 1
            ws.row_dimensions[r].hidden = True
            for c in ws[r]:
                c.fill = GRAU
        elif z["status"] not in EXPORT_STATUS:
            ws.cell(r, SPALTEN.index("Status") + 1).fill = GELB
    ws.sheet_properties.outlinePr.summaryBelow = False
    dv = DataValidation(type="list", formula1='"' + ",".join(URTEILE) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    spalte_urteil = get_column_letter(SPALTEN.index("Urteil") + 1)
    dv.add(f"{spalte_urteil}2:{spalte_urteil}{max(ws.max_row, 2) + 200}")
    for i, name in enumerate(SPALTEN, start=1):
        ws.column_dimensions[get_column_letter(i)].width = BREITEN.get(name, 12)
        if name in ("Pfad", "Lauf"):
            ws.column_dimensions[get_column_letter(i)].hidden = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{ws.max_row}"

    info = wb.create_sheet("Hinweise")
    marker = statistik.get("aufloesung", {}).get("marker", []) if statistik else []
    zeilen = [
        ("Root-Material", root), ("Kurztext", maktx.get(root, "")), ("Lauf", lauf_id),
        ("Stichtag", str(statistik.get("stichtag", "")) if statistik else ""), ("Marker des Laufs", ", ".join(marker)),
        ("", ""),
        ("Urteil", "Bedeutung"),
        ("richtig", "Der Status der Zeile stimmt (drin bei basis/unbedingt, draußen sonst)."),
        ("fehlt", "Das Material gehört in die Basis-Version, ist aber nicht drin. Fehlende Materialien als neue Zeile "
                  "mit Root, Parent, Material, Menge kumuliert ergänzen."),
        ("gehoert_nicht_rein", "Die Zeile ist drin (basis/unbedingt), gehört aber nicht in die Basis-Version."),
        ("", ""),
        ("Ausgeschlossene Zeilen", "stehen eingeklappt am Ende (Gliederung links aufklappen)."),
    ]  # fmt: skip
    if "knart_fehlt" in marker:
        zeilen.append(("ACHTUNG", "CUKB.KNART fehlt im Export: alle Beziehungen wurden als Auswahlbedingungen "
                                  "behandelt, auch Prozeduren (D10). Positionen mit Prozeduren kritisch prüfen."))  # fmt: skip
    for k, v in zeilen:
        info.append([k, v])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 110
    ziel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ziel)
    return ziel


def importiere(eng: Engine, datei: Path, reviewer: str) -> dict:
    """Ausgefüllte Blätter nach `review`; vollständig `richtig` → `bestaetigt`. Liefert eine Zusammenfassung."""
    wb = load_workbook(datei, read_only=True, data_only=True)
    ws = wb[BLATT]
    kopf = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    idx = {name: kopf.index(name) for name in SPALTEN if name in kopf}
    fehlend = [n for n in ("Root", "Parent", "Material", "Urteil") if n not in idx]
    if fehlend:
        raise ValueError(f"{datei.name}: Spalten fehlen: {fehlend}")
    zeilen, fehler, ohne_urteil = [], [], 0
    for nr, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v in (None, "") for v in r):
            continue
        urteil = (r[idx["Urteil"]] or "").strip() if isinstance(r[idx["Urteil"]], str) else r[idx["Urteil"]]
        if not urteil:
            ohne_urteil += 1
            continue
        if urteil not in URTEILE:
            fehler.append(f"Zeile {nr}: Urteil {urteil!r} unbekannt")
            continue

        def feld(name, r=r):
            return r[idx[name]] if name in idx else None

        zeilen.append({
            "root_matnr": str(feld("Root")).strip().lstrip("0"), "matnr": str(feld("Material") or "").strip(),
            "parent_matnr": str(feld("Parent") or "").strip(), "menge": feld("Menge kumuliert"),
            "status": feld("Status"), "urteil": urteil, "kommentar": feld("Kommentar"), "reviewer": reviewer,
            "lauf_id": feld("Lauf"),
        })  # fmt: skip
    wb.close()
    if fehler:
        raise ValueError("; ".join(fehler))
    roots = sorted({z["root_matnr"] for z in zeilen})
    bestaetigt = []
    with eng.begin() as con:
        for z in zeilen:
            con.execute(
                sa.text(
                    "INSERT INTO basis_bom.review (root_matnr, matnr, parent_matnr, menge, status, urteil, kommentar,"
                    " reviewer, lauf_id) VALUES (:root_matnr, :matnr, :parent_matnr, :menge, :status, :urteil,"
                    " :kommentar, :reviewer, :lauf_id)"
                ),
                z,
            )
        for root in roots:
            eigene = [z for z in zeilen if z["root_matnr"] == root]
            if ohne_urteil == 0 and eigene and all(z["urteil"] == "richtig" for z in eigene):
                con.execute(
                    sa.text(
                        "INSERT INTO basis_bom.bestaetigt (root_matnr, bestaetigt_von) VALUES (:r, :v) "
                        "ON CONFLICT (root_matnr) DO UPDATE SET bestaetigt_von = :v, datum = current_date"
                    ),
                    {"r": root, "v": reviewer},
                )
                bestaetigt.append(root)
    return {"zeilen": len(zeilen), "ohne_urteil": ohne_urteil, "roots": roots, "bestaetigt": bestaetigt}
